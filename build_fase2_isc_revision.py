from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import RGBColor

SOURCE = r"C:\Users\brand\OneDrive\Documentos\ChatGPT\chatgpt\campustrade\entregables\CampusTrade_Fase2_Equipo_3_MEX_PER.docx"
OUTPUT = r"C:\Users\brand\OneDrive\Documentos\ChatGPT\chatgpt\campustrade\entregables\CampusTrade_Fase2_Sistemas_ISC_Equipo_3_MEX_PER.docx"

DARK = "173C33"
MINT = "E8F5EE"
PALE = "F7F9F7"
GRAY = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        from docx.oxml import OxmlElement
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = 0
    run = p.add_run(text)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = cell.paragraphs[0].style.font.size
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph(paragraph, replacements):
    text = paragraph.text
    for old, new in replacements.items():
        text = text.replace(old, new)
    if text != paragraph.text:
        paragraph.text = ""
        run = paragraph.add_run(text)
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")


def remove_row(table, index):
    table._tbl.remove(table.rows[index]._tr)


doc = Document(SOURCE)

# Portada y contexto. El documento sigue siendo del equipo común, pero su alcance queda delimitado a ISC.
doc.paragraphs[1].text = "Fase 2 de Sistemas Computacionales"
for run in doc.paragraphs[1].runs:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

meta = doc.tables[0]
set_text(meta.rows[0].cells[1], "Equipo 3 MEX-PER: 4 integrantes. Este documento corresponde únicamente a los 2 integrantes de ISC.")
set_text(meta.rows[1].cells[1], "CampusTrade. Entregable técnico de Sistemas Computacionales.")
set_text(meta.rows[2].cells[1], "Requerimientos, Gantt, comunicación, cambios y pruebas del prototipo alfa.")
set_text(meta.rows[3].cells[1], "Versión alfa sin base de datos. Alcance exclusivo de ISC.")

replacements = {
    "Este documento integra la validación de requerimientos, el seguimiento de actividades, el registro de comunicación, el historial de cambios y las pruebas del prototipo de CampusTrade. La evidencia de comunicación deberá anexarse únicamente con capturas, enlaces o archivos reales del equipo.": "Este documento integra únicamente los entregables de Sistemas Computacionales para CampusTrade. El equipo general tiene cuatro integrantes, pero este alcance corresponde a dos integrantes de ISC. La evidencia se anexará solo con capturas reales de mensajes de WhatsApp y reuniones de Google Meet.",
    "Los requerimientos se revisaron con el anteproyecto, los resultados disponibles de la encuesta aplicada a 208 estudiantes y el alcance del prototipo alfa. La muestra se integra por 70 estudiantes de Industrias Alimentarias, 67 de Mecatrónica y 71 de Automotriz.": "Los dos integrantes de ISC validaron los requerimientos con el anteproyecto, los resultados disponibles de la encuesta aplicada a 208 estudiantes y el alcance del prototipo alfa. Este análisis se utiliza únicamente como fundamento técnico para la plataforma.",
    "El siguiente Gantt resume las actividades del proyecto. Las semanas funcionan como referencia de seguimiento; el equipo puede reemplazarlas por las fechas oficiales del curso si las tiene disponibles.": "El siguiente Gantt resume únicamente las actividades de Sistemas Computacionales. Las semanas son una referencia de seguimiento para los dos integrantes de ISC.",
    "Cada actividad debe actualizarse con su evidencia real: enlace de repositorio, captura del prototipo, documento revisado o registro de comunicación. El estado Completado se utiliza cuando la evidencia está disponible y revisada por el equipo.": "Cada actividad de ISC debe actualizarse con captura del prototipo, documento técnico o comunicación real mediante WhatsApp o Google Meet. El estado Completado se utiliza cuando la evidencia está disponible y fue revisada por los dos integrantes de ISC.",
    "Las evidencias no se sustituyen por conversaciones inventadas. En los espacios siguientes se deben agregar capturas auténticas con fecha visible, enlaces a archivos compartidos o referencias del repositorio del Equipo 3 MEX-PER.": "Las evidencias no se sustituyen por conversaciones inventadas. En los espacios siguientes se deben anexar solo capturas auténticas de mensajes de WhatsApp o reuniones de Google Meet relacionadas con las tareas de ISC.",
    "El historial registra ajustes derivados del anteproyecto, la encuesta y la construcción del prototipo. Los cambios favorecen la reutilización de materiales académicos y la confianza entre estudiantes.": "El historial registra únicamente ajustes técnicos a los requerimientos de CampusTrade realizados por Sistemas Computacionales.",
    "Las pruebas se aplican al prototipo web de CampusTrade. Los resultados Aprobada corresponden a funciones visibles y navegables de la versión alfa. Los elementos Pendiente pertenecen a módulos que requieren autenticación, API o base de datos en una siguiente fase.": "Las pruebas se aplican al prototipo web de CampusTrade elaborado por ISC. Los resultados Aprobada corresponden a funciones visibles y navegables de la versión alfa. Los elementos Pendiente requieren autenticación, API o base de datos en una siguiente fase.",
}
for paragraph in doc.paragraphs:
    replace_paragraph(paragraph, replacements)

# Gantt: solo actividades y responsables ISC.
gantt = doc.tables[4]
activities = [
    "Revisión de requerimientos y datos de encuesta",
    "Validación de requerimientos técnicos",
    "Diseño de interfaz y arquitectura React",
    "Construcción del prototipo alfa",
    "Pruebas funcionales del prototipo",
    "Ajustes visuales y de navegación",
    "Evidencias de WhatsApp y Google Meet",
    "Integración del entregable de ISC",
]
for index, activity in enumerate(activities, start=1):
    set_text(gantt.rows[index].cells[1], activity)
    set_text(gantt.rows[index].cells[2], "ISC", align=WD_ALIGN_PARAGRAPH.CENTER)

# Herramientas de comunicación: únicamente WhatsApp y Google Meet.
tools = doc.tables[5]
set_text(tools.rows[1].cells[0], "WhatsApp")
set_text(tools.rows[1].cells[1], "Acordar alcance técnico, repartir tareas ISC, revisar requerimientos y confirmar ajustes.")
set_text(tools.rows[1].cells[2], "Captura del chat con fecha, hora y contexto del acuerdo técnico.")
set_text(tools.rows[2].cells[0], "Google Meet")
set_text(tools.rows[2].cells[1], "Revisar el prototipo, ejecutar pruebas y acordar correcciones entre los dos integrantes de ISC.")
set_text(tools.rows[2].cells[2], "Captura de la reunión con participantes y pantalla del prototipo o minuta.")
remove_row(tools, 4)
remove_row(tools, 3)

# Guía concreta de capturas reales.
evidence = doc.tables[6]
headers = ["ID", "Medio y tema", "Qué debe verse en la captura", "Uso en el entregable", "Estado"]
for idx, header in enumerate(headers):
    set_text(evidence.rows[0].cells[idx], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color="FFFFFF")
    shade(evidence.rows[0].cells[idx], DARK)
guide = [
    ["EC-01", "WhatsApp: alcance técnico", "Nombre o participantes del chat, fecha y mensaje donde ISC acuerda la función de CampusTrade.", "Evidencia de definición del alcance.", "Por anexar"],
    ["EC-02", "WhatsApp: reparto ISC", "Fecha y mensajes donde los dos integrantes de ISC se distribuyen interfaz, lógica, pruebas o documentación.", "Evidencia de asignación de tareas.", "Por anexar"],
    ["EC-03", "WhatsApp: requerimientos", "Mensajes sobre filtros, búsqueda, modalidades, verificación o ajustes al prototipo.", "Evidencia de validación y cambios.", "Por anexar"],
    ["EC-04", "Google Meet: revisión", "Ventana de Meet con los participantes y, de ser posible, la pantalla compartida del prototipo o lista de pruebas.", "Evidencia de revisión y pruebas.", "Por anexar"],
    ["EC-05", "WhatsApp: cierre ISC", "Mensaje con pruebas revisadas, correcciones acordadas o confirmación del entregable técnico.", "Evidencia de revisión final.", "Por anexar"],
]
for row_index, values in enumerate(guide, start=1):
    for cell_index, value in enumerate(values):
        set_text(evidence.rows[row_index].cells[cell_index], value, align=WD_ALIGN_PARAGRAPH.CENTER if cell_index in {0, 4} else WD_ALIGN_PARAGRAPH.LEFT)
        if row_index % 2 == 0:
            shade(evidence.rows[row_index].cells[cell_index], PALE)

minutes = doc.tables[7]
set_text(minutes.rows[2].cells[1], "Integrantes ISC: ____________________________")
set_text(minutes.rows[5].cells[1], "Captura de Google Meet o registro del acuerdo en WhatsApp")

# Historial y trazabilidad: responsables estrictamente ISC.
changes = doc.tables[8]
for row in changes.rows[1:]:
    set_text(row.cells[4], "ISC", align=WD_ALIGN_PARAGRAPH.CENTER)
set_text(changes.rows[8].cells[2], "Se separaron interfaz, datos simulados y futura API.")
trace = doc.tables[9]
set_text(trace.rows[3].cells[2], "Próxima fase técnica: autenticación, reportes y gestión del sistema")
set_text(trace.rows[4].cells[2], "Próxima fase: autenticación, reportes y gestión del sistema")

# Se evitan referencias al área de Administración: RF-14 es una función futura del sistema.
requirements = doc.tables[2]
set_text(requirements.rows[14].cells[1], "Gestión del sistema")
tests = doc.tables[10]
set_text(tests.rows[12].cells[2], "Registrar reporte para gestión del sistema.")
set_text(tests.rows[12].cells[4], "Requiere módulo de gestión y base de datos.")

# Cierre de pruebas: enfatizar alcance técnico.
for paragraph in doc.paragraphs:
    replace_paragraph(paragraph, {
        "La versión alfa valida la navegación principal, el catálogo, la búsqueda, los filtros, el detalle, los favoritos, la mensajería simulada, la publicación local y la adaptación a pantallas pequeñas.": "La versión alfa de ISC valida la navegación principal, el catálogo, la búsqueda, los filtros, el detalle, los favoritos, la mensajería simulada, la publicación local y la adaptación a pantallas pequeñas.",
        "La versión alfa de ISC valida la navegación principal, el catálogo, la búsqueda, los filtros, el detalle, los favoritos, la mensajería simulada, la publicación local y la adaptación a pantallas pequeñas. La siguiente fase debe incorporar autenticación institucional, almacenamiento persistente, imágenes, operaciones, calificaciones, reportes y administración.": "La versión alfa de ISC valida la navegación principal, el catálogo, la búsqueda, los filtros, el detalle, los favoritos, la mensajería simulada, la publicación local y la adaptación a pantallas pequeñas. La siguiente fase debe incorporar autenticación institucional, almacenamiento persistente, imágenes, operaciones, calificaciones, reportes y gestión del sistema.",
    })

doc.save(OUTPUT)
print(OUTPUT)
