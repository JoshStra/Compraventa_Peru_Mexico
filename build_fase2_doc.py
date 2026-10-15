from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = r"C:\Users\brand\OneDrive\Documentos\ChatGPT\chatgpt\campustrade\entregables\CampusTrade_Fase2_Equipo_3_MEX_PER.docx"

GREEN = "0E8069"
DARK = "173C33"
MINT = "E8F5EE"
PALE = "F7F9F7"
GRAY = "D9D9D9"
ORANGE = "F4B849"
RED = "FBE9E5"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def format_run(run, size=9, bold=False, color=None, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text, size=8.4, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(text))
    format_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(cell)
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths, font_size=8.2, header_fill=DARK):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, header in enumerate(headers):
        cell = header_row.cells[idx]
        set_col_width(cell, widths[idx])
        set_cell_shading(cell, header_fill)
        write_cell(cell, header, size=8.1, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_col_width(cell, widths[idx])
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in {0, len(row) - 1} and len(str(value)) < 28 else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(cell, value, size=font_size, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, number, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(f"{number}. {title}")
    format_run(r, size=15, bold=True, color="000000")
    return p


def add_subheading(doc, title):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    format_run(r, size=11, bold=True, color="000000")
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.13
    if bold_lead:
        r = p.add_run(bold_lead)
        format_run(r, size=10.2, bold=True, color=DARK)
    r = p.add_run(text)
    format_run(r, size=10.2, color=DARK)
    return p


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("CampusTrade  |  Equipo 3 MEX-PER  |  Fase 2")
    format_run(r, size=8, color="60756D")


def status_cell(cell, text):
    palette = {"Aprobada": (MINT, GREEN), "Pendiente": ("FFF4DB", "8A5E0A"), "No aplica": ("EEF1F0", "60756D")}
    fill, color = palette[text]
    set_cell_shading(cell, fill)
    write_cell(cell, text, size=8, bold=True, color=color, align=WD_ALIGN_PARAGRAPH.CENTER)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
add_footer(section)

for style_name in ("Normal", "Heading 1", "Heading 2", "Title"):
    style = doc.styles[style_name]
    style.font.name = "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

# Portada breve
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(62)
p.paragraph_format.space_after = Pt(7)
r = p.add_run("CampusTrade")
format_run(r, size=30, bold=True, color="000000")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
r = p.add_run("Fase 2  Seguimiento y validación")
format_run(r, size=16, bold=False, color="000000")

meta = [
    ("Equipo", "Equipo 3 MEX-PER"),
    ("Proyecto", "Plataforma web de compra, venta, renta, intercambio y donación entre estudiantes"),
    ("Alcance inicial", "Automotriz, Industrias Alimentarias y Mecatrónica"),
    ("Versión", "Prototipo alfa sin base de datos"),
]
table = doc.add_table(rows=0, cols=2)
table.autofit = False
for label, value in meta:
    cells = table.add_row().cells
    set_col_width(cells[0], 1.55)
    set_col_width(cells[1], 4.85)
    set_cell_shading(cells[0], MINT)
    write_cell(cells[0], label, size=9.5, bold=True, color=DARK)
    write_cell(cells[1], value, size=9.5, color=DARK)

doc.add_paragraph().paragraph_format.space_after = Pt(18)
add_body(doc, "Este documento integra la validación de requerimientos, el seguimiento de actividades, el registro de comunicación, el historial de cambios y las pruebas del prototipo de CampusTrade. La evidencia de comunicación deberá anexarse únicamente con capturas, enlaces o archivos reales del equipo.")

# 1 Requirements
add_heading(doc, "1", "Documento de requerimientos validado")
add_body(doc, "Los requerimientos se revisaron con el anteproyecto, los resultados disponibles de la encuesta aplicada a 208 estudiantes y el alcance del prototipo alfa. La muestra se integra por 70 estudiantes de Industrias Alimentarias, 67 de Mecatrónica y 71 de Automotriz.")
add_subheading(doc, "Criterios de validación")
criteria = [
    ["Necesidad", "Los materiales específicos y el uso único de recursos justifican un catálogo organizado."],
    ["Confianza", "El correo institucional y la reputación responden a la necesidad de validar a los participantes."],
    ["Entrega", "La opción preferida es la entrega presencial dentro de la institución."],
    ["Viabilidad alfa", "La primera versión valida la experiencia sin pagos, autenticación real ni base de datos."],
]
add_table(doc, ["Criterio", "Resultado"], criteria, [1.45, 5.25], font_size=8.7)
add_subheading(doc, "Requerimientos funcionales")
requirements = [
    ["RF-01", "Registro e inicio de sesión", "Correo institucional, matrícula, contraseña, carrera y semestre.", "Validado"],
    ["RF-02", "Verificación universitaria", "Comprobar que el usuario pertenece a la institución.", "Validado"],
    ["RF-03", "Perfil de estudiante", "Mostrar carrera, semestre, reputación e historial público.", "Validado"],
    ["RF-04", "Publicar artículo", "Título, descripción, imágenes, carrera, categoría, condición, disponibilidad y modalidad.", "Validado"],
    ["RF-05", "Modalidades", "Venta, renta, intercambio y donación.", "Validado"],
    ["RF-06", "Buscador visible", "Buscar por palabras clave en título y descripción.", "Validado"],
    ["RF-07", "Filtros", "Carrera, categoría, condición, modalidad, precio y disponibilidad.", "Validado"],
    ["RF-08", "Detalle del artículo", "Fotos, descripción, condición, disponibilidad y propietario.", "Validado"],
    ["RF-09", "Mensajería", "Chat interno para acordar condiciones y entrega.", "Validado"],
    ["RF-10", "Favoritos", "Guardar y retirar publicaciones de interés.", "Validado"],
    ["RF-11", "Calificaciones", "Estrellas y comentario después de una operación finalizada.", "Validado"],
    ["RF-12", "Reportes", "Fraude, spam, contenido inapropiado o producto prohibido.", "Validado"],
    ["RF-13", "Entrega", "Acordar entrega presencial dentro de la institución.", "Validado"],
    ["RF-14", "Administración", "Revisar reportes, publicaciones, usuarios, categorías y carreras.", "Validado"],
]
add_table(doc, ["ID", "Requerimiento", "Descripción validada", "Estado"], requirements, [0.55, 1.55, 3.85, 0.75], font_size=7.65)
add_subheading(doc, "Requerimientos no funcionales")
nf = [
    ["RNF-01", "Usabilidad", "Interfaz clara, consistente y fácil de recorrer."],
    ["RNF-02", "Responsividad", "Funcionamiento en computadora, tableta y teléfono."],
    ["RNF-03", "Privacidad", "La información sensible no debe mostrarse públicamente."],
    ["RNF-04", "Escalabilidad", "Posibilidad de incorporar carreras y categorías."],
    ["RNF-05", "Mantenibilidad", "Separación entre interfaz, lógica de negocio y datos."],
]
add_table(doc, ["ID", "Atributo", "Criterio"], nf, [0.75, 1.4, 4.55], font_size=8.2)

# 2 Gantt
doc.add_page_break()
add_heading(doc, "2", "Seguimiento del Gantt")
add_body(doc, "El siguiente Gantt resume las actividades del proyecto. Las semanas funcionan como referencia de seguimiento; el equipo puede reemplazarlas por las fechas oficiales del curso si las tiene disponibles.")
gantt_rows = [
    ["1", "Planeación y definición del alcance", "Equipo 3", "■", "■", "", "", "", "", "Completado"],
    ["2", "Encuesta y análisis de resultados", "Equipo 3", "■", "■", "", "", "", "", "Completado"],
    ["3", "Validación de requerimientos", "Equipo 3", "", "■", "■", "", "", "", "Completado"],
    ["4", "Diseño de interfaz y arquitectura", "Área técnica", "", "", "■", "■", "", "", "Completado"],
    ["5", "Construcción del prototipo alfa", "Área técnica", "", "", "", "■", "■", "", "Completado"],
    ["6", "Pruebas y ajustes del prototipo", "Equipo 3", "", "", "", "", "■", "■", "En progreso"],
    ["7", "Evidencias de comunicación", "Equipo 3", "■", "■", "■", "■", "■", "■", "En progreso"],
    ["8", "Integración de entregable final", "Equipo 3", "", "", "", "", "", "■", "Pendiente"],
]
gantt = add_table(doc, ["#", "Actividad", "Responsable", "S1", "S2", "S3", "S4", "S5", "S6", "Estado"], gantt_rows, [0.28, 2.25, 1.05, 0.34, 0.34, 0.34, 0.34, 0.34, 0.34, 0.82], font_size=7.55)
for row in gantt.rows[1:]:
    for cell in row.cells[3:9]:
        if cell.text.strip() == "■":
            set_cell_shading(cell, GREEN)
            write_cell(cell, "", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    state = row.cells[9].text.strip()
    if state == "Completado":
        set_cell_shading(row.cells[9], MINT)
    elif state == "En progreso":
        set_cell_shading(row.cells[9], "FFF4DB")
    else:
        set_cell_shading(row.cells[9], "EEF1F0")
    write_cell(row.cells[9], state, size=7.35, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
add_subheading(doc, "Criterio de seguimiento")
add_body(doc, "Cada actividad debe actualizarse con su evidencia real: enlace de repositorio, captura del prototipo, documento revisado o registro de comunicación. El estado Completado se utiliza cuando la evidencia está disponible y revisada por el equipo.")

# 3 Communication
doc.add_page_break()
add_heading(doc, "3", "Evidencia del uso de herramientas de comunicación")
add_body(doc, "Las evidencias no se sustituyen por conversaciones inventadas. En los espacios siguientes se deben agregar capturas auténticas con fecha visible, enlaces a archivos compartidos o referencias del repositorio del Equipo 3 MEX-PER.")
media = [
    ["WhatsApp o Telegram", "Acuerdos rápidos, distribución de tareas y avisos.", "Captura del chat con fecha y participantes."],
    ["Reunión presencial o videollamada", "Revisión de avances y decisiones.", "Minuta, foto o captura de videollamada."],
    ["Google Drive o Docs", "Edición colaborativa de documentos.", "Enlace, comentarios o historial de versiones."],
    ["Repositorio", "Seguimiento de código y cambios técnicos.", "Commits, issues, ramas o revisiones."],
]
add_table(doc, ["Herramienta", "Uso en el proyecto", "Evidencia que se anexa"], media, [1.55, 2.65, 2.5], font_size=8.3)
add_subheading(doc, "Registro de evidencias")
evidence = [
    ["EC-01", "Definición del alcance", "Equipo 3 MEX-PER", "Pendiente de anexar captura real", "Pendiente"],
    ["EC-02", "Distribución de tareas", "Equipo 3 MEX-PER", "Pendiente de anexar captura real", "Pendiente"],
    ["EC-03", "Encuesta y requerimientos", "Equipo 3 MEX-PER", "Pendiente de anexar captura real", "Pendiente"],
    ["EC-04", "Revisión del prototipo y pruebas", "Equipo 3 MEX-PER", "Pendiente de anexar captura real", "Pendiente"],
    ["EC-05", "Revisión final del entregable", "Equipo 3 MEX-PER", "Pendiente de anexar captura real", "Pendiente"],
]
add_table(doc, ["ID", "Tema", "Participantes", "Referencia o evidencia", "Estado"], evidence, [0.55, 1.55, 1.35, 2.2, 0.8], font_size=8.0)
add_subheading(doc, "Formato de minuta")
minutes = [
    ["Fecha y medio", "____________________________"],
    ["Participantes", "Equipo 3 MEX-PER: ____________________________"],
    ["Objetivo", "____________________________"],
    ["Acuerdos y tareas", "____________________________"],
    ["Responsables y fecha compromiso", "____________________________"],
    ["Enlace o captura", "____________________________"],
]
add_table(doc, ["Campo", "Registro"], minutes, [1.75, 4.95], font_size=8.5)

# 4 History
doc.add_page_break()
add_heading(doc, "4", "Historial de cambios en requerimientos")
add_body(doc, "El historial registra ajustes derivados del anteproyecto, la encuesta y la construcción del prototipo. Los cambios favorecen la reutilización de materiales académicos y la confianza entre estudiantes.")
changes = [
    ["CH-01", "Modalidades", "Se incluyó formalmente la renta junto con venta, intercambio y donación.", "Ampliar opciones de reutilización.", "Equipo 3"],
    ["CH-02", "Registro", "Se añadió matrícula y verificación universitaria.", "Aumentar confianza e identificación.", "Equipo 3"],
    ["CH-03", "Publicación", "Se añadió condición y disponibilidad del artículo.", "Mostrar el estado real del material.", "Área técnica"],
    ["CH-04", "Búsqueda", "Se priorizó barra de búsqueda visible y filtros combinables.", "Localizar materiales específicos con rapidez.", "Área técnica"],
    ["CH-05", "Seguridad", "Se priorizó el correo institucional como mecanismo de confianza.", "Preferencia identificada en la encuesta.", "Equipo 3"],
    ["CH-06", "Entrega", "Se priorizó el acuerdo de entrega presencial dentro del campus.", "Opción más preferida por los participantes.", "Equipo 3"],
    ["CH-07", "Monetización", "No se considera cuota obligatoria para el MVP.", "La mayoría no desea pagar una tarifa.", "Equipo 3"],
    ["CH-08", "Arquitectura", "Se separaron interfaz, repositorio de datos simulado y futura API.", "Preparar evolución hacia base de datos.", "Área técnica"],
]
add_table(doc, ["ID", "Área", "Cambio", "Motivo o impacto", "Responsable"], changes, [0.55, 1.0, 2.25, 2.1, 0.8], font_size=7.75)
add_subheading(doc, "Trazabilidad del prototipo")
trace = [
    ["Catálogo y búsqueda", "RF-04, RF-06, RF-07, RF-08", "Publicaciones, filtros y detalle"],
    ["Interacción", "RF-09, RF-10", "Mensajería simulada y favoritos"],
    ["Confianza", "RF-02, RF-03, RF-11, RF-13", "Verificación visual, reputación y entrega en campus"],
    ["Evolución", "RF-01, RF-12, RF-14", "Próxima fase: autenticación, reportes y administración"],
]
add_table(doc, ["Módulo", "Requerimientos", "Evidencia en prototipo"], trace, [1.5, 2.15, 3.2], font_size=8.25)

# 5 Tests
doc.add_page_break()
add_heading(doc, "5", "Resultados de pruebas al prototipo")
add_body(doc, "Las pruebas se aplican al prototipo web de CampusTrade. Los resultados Aprobada corresponden a funciones visibles y navegables de la versión alfa. Los elementos Pendiente pertenecen a módulos que requieren autenticación, API o base de datos en una siguiente fase.")
tests = [
    ["PT-01", "Registro", "Crear cuenta con datos válidos.", "No aplica", "Pendiente de autenticación real."],
    ["PT-02", "Inicio de sesión", "Validar credenciales.", "No aplica", "Pendiente de autenticación real."],
    ["PT-03", "Publicación", "Crear artículo desde formulario.", "Aprobada", "Se agrega al catálogo local durante la demostración."],
    ["PT-04", "Modalidades", "Mostrar venta, renta, intercambio y donación.", "Aprobada", "Las cuatro modalidades se muestran en catálogo y formulario."],
    ["PT-05", "Buscador", "Buscar por palabra clave.", "Aprobada", "El catálogo reduce resultados según título, descripción o categoría."],
    ["PT-06", "Filtros", "Filtrar por carrera, condición y modalidad.", "Aprobada", "Los filtros se combinan y actualizan resultados."],
    ["PT-07", "Detalle", "Abrir una publicación.", "Aprobada", "Se muestran condición, entrega, propietario y reputación."],
    ["PT-08", "Mensajería", "Abrir conversación y enviar mensaje simulado.", "Aprobada", "La conversación se muestra en la interfaz sin persistencia."],
    ["PT-09", "Favoritos", "Agregar y retirar publicaciones.", "Aprobada", "La lista de favoritos se actualiza en memoria."],
    ["PT-10", "Finalizar operación", "Cambiar a vendido, rentado o donado.", "Pendiente", "Requiere operación persistente y reglas de permisos."],
    ["PT-11", "Calificación", "Evaluar después de una operación.", "Pendiente", "Requiere operación finalizada y almacenamiento."],
    ["PT-12", "Reporte", "Registrar reporte para administración.", "Pendiente", "Requiere módulo administrativo y base de datos."],
    ["PT-13", "Responsividad", "Abrir en escritorio y teléfono.", "Aprobada", "La interfaz se adapta a ambos tamaños de pantalla."],
]
test_table = add_table(doc, ["ID", "Módulo", "Prueba", "Estado", "Resultado observado"], tests, [0.5, 1.0, 2.2, 0.8, 2.2], font_size=7.65)
for row in test_table.rows[1:]:
    status = row.cells[3].text.strip()
    status_cell(row.cells[3], status)
add_subheading(doc, "Resumen de cierre")
add_body(doc, "La versión alfa valida la navegación principal, el catálogo, la búsqueda, los filtros, el detalle, los favoritos, la mensajería simulada, la publicación local y la adaptación a pantallas pequeñas. La siguiente fase debe incorporar autenticación institucional, almacenamiento persistente, imágenes, operaciones, calificaciones, reportes y administración.")
add_subheading(doc, "Evidencia de pruebas por anexar")
test_evidence = [
    ["PT-03 a PT-06", "Capturas de publicación, búsqueda y filtros.", "Pendiente de anexar"],
    ["PT-07 a PT-09", "Capturas de detalle, chat y favoritos.", "Pendiente de anexar"],
    ["PT-13", "Capturas del prototipo en escritorio y teléfono.", "Pendiente de anexar"],
]
add_table(doc, ["Pruebas", "Evidencia requerida", "Estado"], test_evidence, [1.35, 3.85, 1.5], font_size=8.2)

doc.save(OUTPUT)
print(OUTPUT)
